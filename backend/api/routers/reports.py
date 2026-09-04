from datetime import datetime, timedelta
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import Date, and_, cast, func, select
from sqlalchemy.ext.asyncio import AsyncSession
import csv
import io
from openpyxl import Workbook
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.units import inch, cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from api.deps import get_current_active_user
from database_pg import get_db
from models.models_pg import CreditNote, KardexTransaction, Product, Sale, SaleItem, User
from models.user import UserInDB

router = APIRouter()


@router.get("/summary")
async def get_summary(
    db: AsyncSession = Depends(get_db),
    current_user: UserInDB = Depends(get_current_active_user)
) -> Any:
    utc_now = datetime.utcnow()
    local_now = utc_now - timedelta(hours=5)
    local_today_start = local_now.replace(hour=0, minute=0, second=0, microsecond=0)
    utc_today_start = local_today_start + timedelta(hours=5)

    # Ventas de hoy
    result = await db.execute(
        select(
            func.coalesce(func.sum(Sale.total), 0).label("total"),
            func.count(Sale.id).label("count")
        ).where(Sale.date >= utc_today_start)
    )
    sales_data = result.one()

    # Total productos
    result = await db.execute(select(func.count(Product.id)))
    total_products = result.scalar()

    # Productos con bajo stock
    result = await db.execute(
        select(func.count(Product.id)).where(Product.stock <= Product.min_stock_alert)
    )
    low_stock_count = result.scalar()

    # Total empleados
    result = await db.execute(
        select(func.count(User.id)).where(User.role == "EMPLOYEE")
    )
    total_employees = result.scalar()

    # Top 5 productos mas vendidos
    result = await db.execute(
        select(
            SaleItem.name,
            func.sum(SaleItem.quantity).label("quantity"),
            func.sum(SaleItem.quantity * SaleItem.unit_price).label("revenue")
        )
        .group_by(SaleItem.name)
        .order_by(func.sum(SaleItem.quantity).desc())
        .limit(5)
    )
    top_products = [{"name": r[0], "quantity": r[1], "revenue": r[2]} for r in result.all()]

    return {
        "today_revenue": sales_data.total,
        "today_count": sales_data.count,
        "total_products": total_products,
        "low_stock_count": low_stock_count,
        "total_employees": total_employees,
        "top_products": top_products
    }


@router.get("/sales-chart")
async def get_sales_chart(
    db: AsyncSession = Depends(get_db),
    current_user: UserInDB = Depends(get_current_active_user)
) -> List[Dict]:
    last_7_days = []

    for i in range(6, -1, -1):
        day = (datetime.utcnow() - timedelta(days=i)).replace(hour=0, minute=0, second=0, microsecond=0)
        next_day = day + timedelta(days=1)

        result = await db.execute(
            select(func.coalesce(func.sum(Sale.total), 0))
            .where(and_(Sale.date >= day, Sale.date < next_day))
        )
        total = result.scalar()

        last_7_days.append({
            "name": day.strftime("%a %d"),
            "ventas": total or 0
        })

    return last_7_days


@router.get("/daily-history")
async def get_daily_history(
    db: AsyncSession = Depends(get_db),
    current_user: UserInDB = Depends(get_current_active_user)
) -> List[Dict]:
    # Ventas por dia
    result = await db.execute(
        select(
            cast(Sale.date, Date).label("day"),
            func.sum(Sale.total).label("total"),
            func.count(Sale.id).label("count")
        )
        .group_by(cast(Sale.date, Date))
        .order_by(cast(Sale.date, Date).desc())
        .limit(30)
    )
    sales_data = {str(r.day): {"ventas": r.total, "tickets": r.count} for r in result.all()}

    # Compras por dia (Kardex IN * cost_price)
    result = await db.execute(
        select(
            cast(KardexTransaction.date, Date).label("day"),
            func.sum(KardexTransaction.quantity * Product.cost_price).label("total_purchases")
        )
        .join(Product, KardexTransaction.product_id == Product.id)
        .where(KardexTransaction.transaction_type == "IN")
        .group_by(cast(KardexTransaction.date, Date))
        .order_by(cast(KardexTransaction.date, Date).desc())
        .limit(30)
    )
    purchases_data = {str(r.day): r.total_purchases or 0 for r in result.all()}

    # Combinar
    all_dates = set(list(sales_data.keys()) + list(purchases_data.keys()))
    daily_data = []

    for d in all_dates:
        daily_data.append({
            "date": d,
            "ventas": sales_data.get(d, {}).get("ventas", 0) if isinstance(sales_data.get(d), dict) else 0,
            "compras": purchases_data.get(d, 0),
            "tickets": sales_data.get(d, {}).get("tickets", 0) if isinstance(sales_data.get(d), dict) else 0
        })

    daily_data.sort(key=lambda x: x["date"], reverse=True)
    return daily_data[:30]


@router.get("/daily-details/{date_str}")
async def get_daily_details(
    date_str: str,
    db: AsyncSession = Depends(get_db),
    current_user: UserInDB = Depends(get_current_active_user)
) -> Any:
    try:
        local_date = datetime.strptime(date_str, "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=400, detail="Formato de fecha invalido. Usar YYYY-MM-DD")

    start_utc = datetime.combine(local_date, datetime.min.time()) + timedelta(hours=5)
    end_utc = start_utc + timedelta(days=1)

    # Ventas del dia
    result = await db.execute(
        select(Sale).where(and_(Sale.date >= start_utc, Sale.date < end_utc)).order_by(Sale.date.desc())
    )
    sales = result.scalars().all()

    sales_list = []
    for s in sales:
        sales_list.append({
            "id": str(s.id),
            "client_name": s.client_name,
            "total": s.total,
            "date": s.date.isoformat() if s.date else None
        })

    # Compras del dia
    result = await db.execute(
        select(KardexTransaction, Product.name.label("product_name"), Product.cost_price)
        .join(Product, KardexTransaction.product_id == Product.id)
        .where(
            and_(
                KardexTransaction.transaction_type == "IN",
                KardexTransaction.date >= start_utc,
                KardexTransaction.date < end_utc
            )
        )
    )
    purchases = [
        {
            "product_name": r.product_name,
            "cost_price": r.cost_price,
            "quantity": r.quantity,
            "notes": r.notes,
            "date": r.date.isoformat() if r.date else None
        }
        for r in result.all()
    ]

    return {
        "date": date_str,
        "sales": sales_list,
        "purchases": purchases
    }


@router.get("/customer-history/{customer_dni}")
async def get_customer_history(
    customer_dni: str,
    skip: int = 0,
    limit: int = 20,
    db: AsyncSession = Depends(get_db),
    current_user: UserInDB = Depends(get_current_active_user),
) -> Any:
    """Historial de compras de un cliente."""
    result = await db.execute(
        select(Sale)
        .where(Sale.client_id == customer_dni)
        .order_by(Sale.date.desc())
        .offset(skip)
        .limit(limit)
    )
    sales = result.scalars().all()

    # Total gastado
    total_result = await db.execute(
        select(func.coalesce(func.sum(Sale.total), 0))
        .where(Sale.client_id == customer_dni)
    )
    total_spent = total_result.scalar()

    # Conteo de compras
    count_result = await db.execute(
        select(func.count(Sale.id))
        .where(Sale.client_id == customer_dni)
    )
    purchase_count = count_result.scalar()

    return {
        "customer_dni": customer_dni,
        "total_spent": total_spent,
        "purchase_count": purchase_count,
        "sales": [
            {
                "id": str(s.id),
                "client_name": s.client_name,
                "subtotal": s.subtotal,
                "tax": s.tax,
                "total": s.total,
                "payment_method": s.payment_method,
                "date": s.date.isoformat() if s.date else None,
                "items_count": len(s.sale_items) if hasattr(s, 'sale_items') else 0,
            }
            for s in sales
        ],
    }


@router.get("/top-products")
async def get_top_products(
    days: int = Query(30, ge=1, le=365),
    limit: int = Query(10, ge=1, le=50),
    db: AsyncSession = Depends(get_db),
    current_user: UserInDB = Depends(get_current_active_user),
) -> Any:
    """Productos mas vendidos por cantidad y reveneu."""
    since = datetime.utcnow() - timedelta(days=days)

    result = await db.execute(
        select(
            SaleItem.name,
            func.sum(SaleItem.quantity).label("total_qty"),
            func.sum(SaleItem.quantity * SaleItem.unit_price).label("total_revenue"),
            func.count(SaleItem.id).label("sales_count"),
        )
        .join(Sale, SaleItem.sale_id == Sale.id)
        .where(Sale.date >= since)
        .group_by(SaleItem.name)
        .order_by(func.sum(SaleItem.quantity).desc())
        .limit(limit)
    )
    products = [
        {"name": r[0], "quantity": r[1], "revenue": r[2], "sales_count": r[3]}
        for r in result.all()
    ]

    return {"days": days, "products": products}


@router.get("/sales-by-method")
async def get_sales_by_method(
    days: int = Query(30, ge=1, le=365),
    db: AsyncSession = Depends(get_db),
    current_user: UserInDB = Depends(get_current_active_user),
) -> Any:
    """Ventas agrupadas por metodo de pago."""
    since = datetime.utcnow() - timedelta(days=days)

    result = await db.execute(
        select(
            Sale.payment_method,
            func.coalesce(func.sum(Sale.total), 0).label("total"),
            func.count(Sale.id).label("count"),
        )
        .where(Sale.date >= since)
        .group_by(Sale.payment_method)
        .order_by(func.sum(Sale.total).desc())
    )
    methods = [
        {"method": r[0] or "cash", "total": r[1], "count": r[2]}
        for r in result.all()
    ]

    return {"days": days, "methods": methods}


@router.get("/credit-notes-summary")
async def get_credit_notes_summary(
    days: int = Query(30, ge=1, le=365),
    db: AsyncSession = Depends(get_db),
    current_user: UserInDB = Depends(get_current_active_user),
) -> Any:
    """Resumen de notas de credito emitidas."""
    since = datetime.utcnow() - timedelta(days=days)

    result = await db.execute(
        select(
            func.coalesce(func.sum(CreditNote.total), 0).label("total"),
            func.count(CreditNote.id).label("count"),
        )
        .where(CreditNote.created_at >= since)
    )
    data = result.one()

    return {"days": days, "total_amount": data.total, "note_count": data.count}


# ===== EXPORTACIONES =====

@router.get("/export/sales")
async def export_sales_csv(
    days: int = Query(30, ge=1, le=365),
    format: str = Query("csv", regex="^(csv|excel|pdf)$"),
    db: AsyncSession = Depends(get_db),
    current_user: UserInDB = Depends(get_current_active_user),
) -> Any:
    """Exportar ventas en CSV, Excel o PDF."""
    since = datetime.utcnow() - timedelta(days=days)

    result = await db.execute(
        select(Sale)
        .where(Sale.date >= since)
        .order_by(Sale.date.desc())
    )
    sales = result.scalars().all()

    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Fecha", "Cliente", "Subtotal", "IVA", "Total", "Metodo Pago", "Estado SRI"])
        for s in sales:
            writer.writerow([
                s.date.strftime("%Y-%m-%d %H:%M") if s.date else "",
                s.client_name,
                f"{s.subtotal:.2f}",
                f"{s.tax:.2f}",
                f"{s.total:.2f}",
                s.payment_method or "cash",
                s.sri_status or "local"
            ])
        output.seek(0)
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=ventas_{days}dias.csv"}
        )

    elif format == "excel":
        wb = Workbook()
        ws = wb.active
        ws.title = "Ventas"
        ws.append(["Fecha", "Cliente", "Subtotal", "IVA", "Total", "Metodo Pago", "Estado SRI"])
        for s in sales:
            ws.append([
                s.date.strftime("%Y-%m-%d %H:%M") if s.date else "",
                s.client_name,
                s.subtotal,
                s.tax,
                s.total,
                s.payment_method or "cash",
                s.sri_status or "local"
            ])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=ventas_{days}dias.xlsx"}
        )

    elif format == "pdf":
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=A4, topMargin=1*cm, bottomMargin=1*cm)
        styles = getSampleStyleSheet()
        elements = []

        title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=16, spaceAfter=20)
        elements.append(Paragraph(f"Reporte de Ventas - Ultimos {days} dias", title_style))
        elements.append(Paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
        elements.append(Spacer(1, 20))

        table_data = [["Fecha", "Cliente", "Subtotal", "IVA", "Total", "Metodo"]]
        for s in sales:
            table_data.append([
                s.date.strftime("%d/%m/%Y") if s.date else "",
                s.client_name[:20],
                f"${s.subtotal:.2f}",
                f"${s.tax:.2f}",
                f"${s.total:.2f}",
                s.payment_method or "Efectivo"
            ])

        table = Table(table_data, colWidths=[2.5*cm, 4*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3b82f6')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f4f6')]),
        ]))
        elements.append(table)

        doc.build(elements)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename=ventas_{days}dias.pdf"}
        )


@router.get("/export/products")
async def export_products_csv(
    format: str = Query("csv", regex="^(csv|excel|pdf)$"),
    db: AsyncSession = Depends(get_db),
    current_user: UserInDB = Depends(get_current_active_user),
) -> Any:
    """Exportar inventario en CSV, Excel o PDF."""
    result = await db.execute(select(Product).order_by(Product.name))
    products = result.scalars().all()

    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Nombre", "Codigo", "Precio Costo", "Precio Venta", "Stock", "Stock Minimo", "IVA"])
        for p in products:
            writer.writerow([
                p.name,
                p.barcode or "",
                f"{p.cost_price:.2f}",
                f"{p.sale_price:.2f}",
                p.stock,
                p.min_stock_alert,
                p.iva_rate or 15
            ])
        output.seek(0)
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=inventario.csv"}
        )

    elif format == "excel":
        wb = Workbook()
        ws = wb.active
        ws.title = "Inventario"
        ws.append(["Nombre", "Codigo", "Precio Costo", "Precio Venta", "Stock", "Stock Minimo", "IVA"])
        for p in products:
            ws.append([p.name, p.barcode or "", p.cost_price, p.sale_price, p.stock, p.min_stock_alert, p.iva_rate or 15])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=inventario.xlsx"}
        )

    elif format == "pdf":
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=A4, topMargin=1*cm, bottomMargin=1*cm)
        styles = getSampleStyleSheet()
        elements = []

        title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=16, spaceAfter=20)
        elements.append(Paragraph("Reporte de Inventario", title_style))
        elements.append(Paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
        elements.append(Spacer(1, 20))

        table_data = [["Producto", "Codigo", "Costo", "Venta", "Stock", "Min"]]
        for p in products:
            table_data.append([
                p.name[:25],
                p.barcode or "-",
                f"${p.cost_price:.2f}",
                f"${p.sale_price:.2f}",
                str(p.stock),
                str(p.min_stock_alert)
            ])

        table = Table(table_data, colWidths=[5*cm, 3*cm, 2.5*cm, 2.5*cm, 2*cm, 2*cm])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#10b981')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f4f6')]),
        ]))
        elements.append(table)

        doc.build(elements)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=inventario.pdf"}
        )


@router.get("/export/top-products")
async def export_top_products_csv(
    days: int = Query(30, ge=1, le=365),
    format: str = Query("csv", regex="^(csv|excel|pdf)$"),
    db: AsyncSession = Depends(get_db),
    current_user: UserInDB = Depends(get_current_active_user),
) -> Any:
    """Exportar top productos."""
    since = datetime.utcnow() - timedelta(days=days)

    result = await db.execute(
        select(
            SaleItem.name,
            func.sum(SaleItem.quantity).label("total_qty"),
            func.sum(SaleItem.quantity * SaleItem.unit_price).label("total_revenue"),
        )
        .join(Sale, SaleItem.sale_id == Sale.id)
        .where(Sale.date >= since)
        .group_by(SaleItem.name)
        .order_by(func.sum(SaleItem.quantity).desc())
        .limit(50)
    )
    products = [{"name": r[0], "quantity": r[1], "revenue": r[2]} for r in result.all()]

    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Producto", "Unidades Vendidas", "Ingresos"])
        for p in products:
            writer.writerow([p["name"], p["quantity"], f"{p['revenue']:.2f}"])
        output.seek(0)
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=top_productos_{days}dias.csv"}
        )

    elif format == "excel":
        wb = Workbook()
        ws = wb.active
        ws.title = "Top Productos"
        ws.append(["Producto", "Unidades Vendidas", "Ingresos"])
        for p in products:
            ws.append([p["name"], p["quantity"], p["revenue"]])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=top_productos_{days}dias.xlsx"}
        )

    elif format == "pdf":
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=A4, topMargin=1*cm, bottomMargin=1*cm)
        styles = getSampleStyleSheet()
        elements = []

        title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=16, spaceAfter=20)
        elements.append(Paragraph(f"Top Productos - Ultimos {days} dias", title_style))
        elements.append(Spacer(1, 20))

        table_data = [["#", "Producto", "Unidades", "Ingresos"]]
        for i, p in enumerate(products, 1):
            table_data.append([str(i), p["name"][:30], str(p["quantity"]), f"${p['revenue']:.2f}"])

        table = Table(table_data, colWidths=[1*cm, 7*cm, 2.5*cm, 3*cm])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#8b5cf6')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f4f6')]),
        ]))
        elements.append(table)

        doc.build(elements)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename=top_productos_{days}dias.pdf"}
        )
